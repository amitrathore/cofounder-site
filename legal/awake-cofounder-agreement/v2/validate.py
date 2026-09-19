#!/usr/bin/env python3
"""Structural/content checks; complements manual legal-scenario and rendered-page review."""
from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document
from docx.oxml.ns import qn
HERE=Path(__file__).resolve().parent
STEMS=['awake-cofounder-agreement-v2-review','guide','change-summary']
s=(HERE/(STEMS[0]+'.md')).read_text()
clauses=re.findall(r'^\*\*(\d+\.\d+)\*\*',s,re.M)
assert len(clauses)==len(set(clauses)), 'Duplicate clause numbers'
for ref in re.findall(r'Clause (\d+\.\d+)',s):assert ref in clauses,ref
for n in range(1,22):assert re.search(r'^## '+str(n)+r'\. ',s,re.M),n
for x in ['Schedules A–D','Original Founder Block','Later Participant','Pro Rata Tax Advances','Tax Classification, QSBS','no personal repayment obligation','shall not supply missing consent','original Member’s responsibility does not establish continuous tax ownership']:
 assert x in s,x
assert 'shall not, for [COUNSEL:' not in s
assert 'full acceleration on a Sale' not in s
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for stem in STEMS:
 d=Document(HERE/(stem+'.docx'))
 sec=d.sections[0]
 assert sec.page_width.twips==12240 and sec.page_height.twips==15840
 assert all(x.twips==1440 for x in [sec.top_margin,sec.bottom_margin,sec.left_margin,sec.right_margin])
 assert d.styles['Normal'].font.size.pt==11
 assert not d.styles['Title'].element.findall('./'+qn('w:pPr')+'/'+qn('w:pBdr'))
 assert 'REVIEW DRAFT' in sec.header.paragraphs[0].text
 doc_text='\n'.join(''.join(('\n' if t.tag==qn('w:br') else (t.text or '')) for t in p._p.iter() if t.tag in [qn('w:t'),qn('w:br')]) for p in d.paragraphs)
 # Compare text after removing rendering whitespace, running furniture, and auto-list markers.
 pdf=PdfReader(HERE/(stem+'.pdf'))
 pieces=[]
 for pg in pdf.pages:
  txt=pg.extract_text()
  txt=re.sub(r'^.*\| V2\.0 PUBLIC REVIEW DRAFT\s*','',txt,flags=re.M)
  txt=re.sub(r'Not approved for signature\s*\|\s*\d+','',txt)
  txt=re.sub(r'(?m)^[ \t]*(?:•|\d{1,2}\.)[ \t]+(?=\S)', '',txt)
  pieces.append(txt)
 # Remove only paragraph-leading list/heading markers from the Word body too.
 doc_text=re.sub(r'(?m)^[ \t]*(?:•|\d{1,2}\.)[ \t]+(?=\S)', '',doc_text)
 def tokens(t):
  t=t.replace('\u00ad','').replace('ﬁ','fi').replace('ﬂ','fl')
  return re.findall(r'[A-Za-z0-9]+',t)
 a=tokens(doc_text);b=tokens('\n'.join(pieces))
 assert a==b, f'{stem}: PDF word sequence mismatch near '+str(next((i for i,(x,y) in enumerate(zip(a,b)) if x!=y), min(len(a),len(b))))
 print(stem, len(pdf.pages), 'pages: Word/PDF word sequence and layout tokens passed')
print('Clause references, numbering, review labels, and protected-structure checks passed.')
