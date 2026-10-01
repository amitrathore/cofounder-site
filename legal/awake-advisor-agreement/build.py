#!/usr/bin/env python3
"""Build the advisor agreement's Word and online editions from canonical Markdown."""
import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
STEM = "awake-advisor-agreement"
blocks = (HERE / f"{STEM}.md").read_text().strip().split("\n\n")
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

for name, size, before, after, color in [
    ("Normal", 11, 0, 6, "161411"),
    ("Title", 24, 0, 8, "161411"),
    ("Subtitle", 11, 0, 6, "555555"),
    ("Heading 1", 16, 16, 8, "B83A2E"),
    ("Heading 2", 13, 12, 6, "B83A2E"),
    ("Header", 9, 0, 0, "555555"),
    ("Footer", 9, 0, 0, "555555"),
]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = name.startswith("Heading")
    pf = style.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.10
    pf.widow_control = True
    pf.keep_with_next = name.startswith("Heading") or name == "Title"
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for border in style.element.findall("./" + qn("w:pPr") + "/" + qn("w:pBdr")):
        border.getparent().remove(border)

section.header.paragraphs[0].text = "AWAKE ADVISOR AGREEMENT  |  VERSION 1.0"
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("cofounder.community  |  CC BY 4.0  |  ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "Awake Advisor Agreement"
doc.core_properties.author = "Amit Rathore and AwakeVC"
doc.core_properties.subject = "Advisor services, common share vesting, and commercial referral commission template"


def add_runs(paragraph, value):
    for part in re.split(r"(\[[^\]]+\])", value):
        run = paragraph.add_run(part)
        if part.startswith("[") and part.endswith("]"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def inline(value):
    return re.sub(r"\[([^\]]+)\]", r"<mark>[\1]</mark>", html.escape(value))


body, toc = [], []
signatures = False
for i, block in enumerate(blocks):
    if block.startswith("# "):
        doc.add_paragraph(block[2:], "Title")
        continue
    match = re.match(r"^(#{2,3}) (.+)$", block)
    if match:
        level, title = len(match[1]), match[2]
        anchor = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")
        p = doc.add_paragraph(title, f"Heading {level - 1}")
        if title == "1. Engagement and services":
            p.paragraph_format.page_break_before = True
        signatures = title == "Signatures"
        body.append(f'<h{level} id="{anchor}">{html.escape(title)}</h{level}>')
        if level == 2:
            toc.append(f'<li><a href="#{anchor}">{html.escape(title)}</a></li>')
    else:
        if "\n" in block or block.startswith(("- ", "* ", "|", "```")):
            raise ValueError("Unsupported Markdown block: " + block[:80])
        p = doc.add_paragraph(style="Subtitle" if i == 1 else "Normal")
        add_runs(p, block)
        if signatures:
            p.paragraph_format.keep_with_next = True
        body.append(f"<p>{inline(block)}</p>")
doc.paragraphs[-1].paragraph_format.keep_with_next = False
doc.save(HERE / f"{STEM}.docx")

page = '''<!DOCTYPE html>
<html lang="en"><head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Awake Advisor Agreement · Version 1.0</title>
<meta name="description" content="Open advisor agreement template with common shares vesting quarterly over three years, strategy services, introductions, and commercial referral commissions.">
<link rel="canonical" href="https://cofounder.community/legal/awake-advisor-agreement/">
<meta property="og:type" content="website"><meta property="og:site_name" content="Cofounder">
<meta property="og:title" content="Awake Advisor Agreement · Version 1.0">
<meta property="og:description" content="Advisor services, quarterly common share vesting, and commercial referral commissions.">
<meta property="og:url" content="https://cofounder.community/legal/awake-advisor-agreement/">
<meta property="og:image" content="https://cofounder.community/assets/og-cofounder-v2.png">
<meta name="twitter:card" content="summary_large_image"><meta name="theme-color" content="#FAF7F2">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Fraunces:opsz,wght@9..144,400..600&amp;family=Inter:wght@400;500;600&amp;display=swap" rel="stylesheet">
<link rel="stylesheet" href="../../assets/book.css"><link rel="stylesheet" href="../../assets/program.css">
<link rel="stylesheet" href="../../assets/docs.css"><script src="../../assets/docs-nav.js" defer></script>
<link rel="stylesheet" href="./reader.css"></head><body>
<nav class="book-nav" aria-label="Main navigation"><div class="book-nav-inner">
<a href="/" class="book-nav-brand">Cofounder<span class="dot">.</span></a>
<div class="book-nav-links"><a href="/legal/">Agreements</a><details class="docs-nav"><summary aria-expanded="false">Docs <span aria-hidden="true">▾</span></summary><div class="docs-dropdown"><a href="/docs/"><strong>All Docs</strong><span>Guides and open templates</span></a><a href="/legal/" aria-current="page"><strong>Agreements</strong><span>Founder and venture relationships</span></a><a href="/docs/data-room/"><strong>Data Room</strong><span>Prepare for investor diligence</span></a></div></details><a href="./awake-advisor-agreement.docx" download>Word</a><a href="./awake-advisor-agreement.pdf" download>PDF</a><a href="./awake-advisor-agreement.md">Source</a><a href="https://github.com/amitrathore/cofounder-site/issues/new">Improve ↗</a></div>
</div></nav>
<main class="legal-reader advisor-reader"><header>
<div class="program-eyebrow">Open business template · Version 1.0</div>
<h1>Awake Advisor Agreement<span class="dot">.</span></h1>
<p>Set out an advisor's strategy and introduction services, a fixed common share award vesting quarterly over three years, and participation in eligible sales commissions.</p>
<div class="program-actions"><a class="program-button primary" href="./awake-advisor-agreement.docx" download>Download editable</a><a class="program-button" href="./awake-advisor-agreement.pdf" download>Download PDF</a><a class="program-button" href="./guide.md">Completion guide</a></div>
</header>
<aside class="legal-alert"><strong>Not legal advice.</strong> Complete the share, commission, referral, and governing-law fields with independent counsel and tax advice. Referral commissions here cover lawful commercial sales, not financing or securities transactions. <a href="./DISCLAIMER.md">Full disclaimer →</a></aside>
<section class="legal-section" aria-labelledby="using-heading"><h2 id="using-heading">Make the economics explicit.</h2>
<p>The equity award uses a fixed number of common shares and twelve quarterly vesting dates. Separate signed equity documents and required corporate approvals implement the grant.</p>
<p>Record each accepted introduction in writing. Specify the commission base, percentage, protection period, and post-termination tail before signing. <a href="./guide.md">The completion guide</a> explains these fields and the regulatory boundary.</p>
<p>Highlighted brackets identify fields to complete. The Word and PDF copies contain the same agreement text. <a href="./awake-advisor-agreement.md">Canonical Markdown</a> · <a href="./CHANGELOG.md">Version history</a> · <a href="./CONTRIBUTING.md">Contribute</a></p></section>
<details class="advisor-contents"><summary>Agreement contents</summary><ul>__TOC__</ul></details>
<article class="advisor-text" aria-label="Awake Advisor Agreement version 1.0">__BODY__</article>
<aside class="legal-contribution"><div class="program-eyebrow">An AwakeVC contribution</div><div class="legal-contribution-copy"><p class="legal-contribution-statement">Contributed by Amit Rathore and AwakeVC for founders, advisors, and entrepreneurial communities.</p><p class="legal-contribution-terms">Free to use, adapt, and share under <a href="./LICENSE.md">CC BY 4.0</a>. Independent legal review: not recorded.</p></div></aside>
</main><footer class="book-foot"><span>&copy; MMXXVI · <a href="/"><strong>cofounder.com</strong>munity</a> · <a href="https://cofounder.exchange">Cofounder.Exchange ↗</a> · <a class="creator-link" href="https://awake.vc">An Awake Venture ↗</a></span></footer>
</body></html>
'''
(HERE / "index.html").write_text(page.replace("__TOC__", "\n".join(toc)).replace("__BODY__", "\n".join(body)))
print("Built online reading page and Word from canonical Markdown.")
