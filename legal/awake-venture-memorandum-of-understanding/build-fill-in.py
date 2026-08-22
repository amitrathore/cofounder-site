#!/usr/bin/env python3
"""Build the fill-in copy of the Awake Venture Memorandum of Understanding.

The fill-in copy is the annotated template with the drafting apparatus removed:
no version block, no licence page, no field legend, no [COUNSEL: ...] notes.
What remains is the operative memorandum with every blank highlighted.

Document content is defined once, below, and emitted to both Word and HTML.
The HTML is rendered to PDF with headless Chromium. Keeping one source means the
two published formats cannot drift apart, which the package's release checks require.

Usage:  python3 build-fill-in.py
"""

import html
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
STEM = "awake-venture-memorandum-of-understanding-fill-in"
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

# Run kinds: "t" plain text, "b" bold, "f" a blank to fill in (highlighted).
T, B, F = "t", "b", "f"

TITLE = "Memorandum of Understanding"

PARTY = [
    [(B, "Venture: "), (F, "venture name"), (T, " (the “Venture”)")],
    [(B, "Partners: "), (F, "partner 1 full legal name"), (T, "; "),
     (F, "partner 2 full legal name"), (T, " (each a “Partner”)")],
    [(B, "Effective date: "), (F, "date"), (T, "   ·   "),
     (B, "Outside date for definitive agreements: "), (F, "date")],
]

BODY = [
    ("p", [(B, "1. Intent. "), (T, "The Partners intend to build "),
           (F, "what the venture makes, for whom, and the problem it solves"),
           (T, ". This memorandum records what they have agreed in principle so that "
               "definitive agreements can be drafted from it.")]),

    ("p", [(B, "2. Structure. "), (T, "The Venture will operate through "), (F, "entity name"),
           (T, ", a "), (F, "entity type and jurisdiction"), (T, " to be formed by "), (F, "date"),
           (T, ". Each Partner will hold their interest "),
           (F, "directly / through a founders’ holding vehicle"),
           (T, ". Initial capital: "), (F, "amount and who contributes it, or “none”"), (T, ".")]),

    ("p", [(B, "3. Equity. "), (T, "Founding equity, fully diluted, before outside investment:")]),
    ("table", {
        "widths": [22, 13, 27, 22, 16],
        "head": ["Partner", "Founding equity", "Consideration", "Vesting", "Cliff"],
        "rows": [
            [[(F, "name")], [(F, "%")], [(F, "cash / services / IP / combination")],
             [(F, "e.g. 4 years, monthly")], [(F, "e.g. 12 months")]],
            [[(F, "name")], [(F, "%")], [(F, "cash / services / IP / combination")],
             [(F, "e.g. 4 years, monthly")], [(F, "e.g. 12 months")]],
            [[(T, "Reserved option pool")], [(F, "%")], [(T, "—")], [(T, "—")], [(T, "—")]],
        ]}),
    ("p", [(T, "Vesting commences "), (F, "date"), (T, ". Acceleration: "),
           (F, "none / single-trigger on change of control / double-trigger on change of control "
               "and involuntary termination"),
           (T, ". Dilution from future financings is borne pro rata by all Partners unless they "
               "unanimously agree otherwise.")]),

    ("p", [(B, "4. Other compensation. "), (T, "Beyond equity, the Partners intend:")]),
    ("table", {
        "widths": [26, 74],
        "head": ["Item", "Terms"],
        "rows": [
            [[(T, "Salary")], [(F, "none until [milestone]; then [amount] per Partner")]],
            [[(T, "Deferred or accrued compensation")],
             [(F, "amount, trigger for payment, and whether it converts to equity, or “none”")]],
            [[(T, "Expense reimbursement")],
             [(F, "what is reimbursable, and any approval threshold")]],
            [[(T, "Profit distributions")],
             [(F, "policy and who approves, or “none until the Partners agree otherwise”")]],
            [[(T, "Other")], [(F, "benefits, advisory fees, contractor rates, or “none”")]],
        ]}),

    ("p", [(B, "5. Roles and responsibilities. "),
           (T, "Each Partner is individually accountable for their column and holds day-to-day "
               "authority within it:")]),
    ("table", {
        "widths": [20, 16, 40, 24],
        "head": ["Partner", "Title", "Accountable for", "Time commitment"],
        "rows": [
            [[(F, "name")], [(F, "title")], [(F, "the 3–5 outcomes this Partner owns")],
             [(F, "e.g. full-time from [date]")]],
            [[(F, "name")], [(F, "title")], [(F, "the 3–5 outcomes this Partner owns")],
             [(F, "e.g. full-time from [date]")]],
        ]}),
    ("p", [(T, "Decisions outside a Partner’s own column require "), (F, "unanimous / __%"),
           (T, " written approval, including: spending above "), (F, "amount"),
           (T, "; hiring; raising capital or incurring debt; issuing equity; changing the business "
               "or this memorandum; entering any contract above "), (F, "amount"),
           (T, "; and selling the Venture. Deadlock: "), (F, "escalation path and time limit"), (T, ".")]),

    ("p", [(B, "6. Intellectual property. "),
           (T, "All work created for the Venture is intended to belong to the entity, and each "
               "Partner will sign an invention-assignment and confidentiality agreement at "
               "formation. Background intellectual property each Partner brings: "),
           (F, "list, or “none”"),
           (T, ", which remains that Partner’s and is licensed to the Venture "),
           (F, "perpetually and royalty-free / on terms to be agreed in the definitive agreements"),
           (T, ".")]),

    ("p", [(B, "7. Exit intent. "),
           (T, "The Partners record the following shared intention and will build toward it:")]),
    ("table", {
        "widths": [26, 74],
        "head": ["Item", "Intent"],
        "rows": [
            [[(T, "Preferred outcome")],
             [(F, "strategic acquisition / initial public offering / independent cash-flow business "
                  "paying distributions / buyout of one or more Partners / no predetermined outcome")]],
            [[(T, "Target window")], [(F, "e.g. 5–7 years from the effective date")]],
            [[(T, "Target valuation or return")],
             [(F, "e.g. not below [amount] enterprise value, or “no target set”")]],
            [[(T, "When a process starts")],
             [(F, "the trigger — e.g. an inbound offer above target, or a Partner vote after [date]")]],
            [[(T, "Approval to accept an offer")], [(F, "unanimous / __% of Partner equity")]],
            [[(T, "Drag-along and tag-along")],
             [(T, "Definitive agreements are intended to include drag-along rights on an approved "
                  "sale and tag-along rights protecting every Partner on a transfer by another.")]],
            [[(T, "If no exit by the target window")],
             [(T, "The Partners will meet within "), (F, "days"), (T, " days and "),
              (F, "extend the window / start a buy-sell at independently determined fair market "
                  "value / wind the Venture down"), (T, ".")]],
            [[(T, "Individual side deals")],
             [(T, "No Partner will negotiate or accept a purchase of their own interest, or a role "
                  "offer conditioned on transferring it, without first disclosing it to the other "
                  "Partners and offering them "),
              (F, "a right of first refusal / the right to participate pro rata"), (T, ".")]],
        ]}),

    ("p", [(B, "8. Definitive agreements. "),
           (T, "The Partners intend to sign, by the outside date above: entity formation documents; "
               "a founders’ or operating agreement; equity purchase documents with the vesting "
               "in Clause 3; invention-assignment and confidentiality agreements; and "),
           (F, "any additional documents, or “none”"),
           (T, ". If those are not signed by the outside date, this memorandum expires unless the "
               "Partners extend it in writing.")]),

    ("p", [(B, "9. Binding effect. "),
           (B, "Clauses 1 through 8 record intent only and are not legally binding. "),
           (T, "No Partner may sue on them, and no obligation to form the Venture, contribute "
               "capital, transfer property, or issue equity arises until definitive agreements are "
               "signed. "),
           (B, "The following are intended to be legally binding and to survive expiry or "
               "termination:")]),
    ("sub", [(B, "9.1 Confidentiality. "),
             (T, "Each Partner will keep the other Partners’ non-public information, and the "
                 "existence and contents of this memorandum, confidential for "), (F, "period"),
             (T, " and use it only to evaluate and pursue the Venture.")]),
    ("sub", [(B, "9.2 Exclusivity. "),
             (T, "Until the outside date, no Partner will pursue, fund, or join a venture that "
                 "competes with the Venture as described in Clause 1, except "),
             (F, "carve-outs, or “no exceptions”"), (T, ".")]),
    ("sub", [(B, "9.3 Costs. "),
             (T, "Each Partner bears their own costs, including legal fees, unless otherwise agreed "
                 "in writing.")]),
    ("sub", [(B, "9.4 No partnership or employment. "),
             (T, "Nothing here creates a partnership, joint venture, agency, fiduciary, or "
                 "employment relationship between the Partners, or authority to bind another "
                 "Partner.")]),
    ("sub", [(B, "9.5 Governing law and disputes. "),
             (T, "This memorandum is governed by the laws of "), (F, "jurisdiction"),
             (T, ". Disputes will be resolved by "),
             (F, "good-faith negotiation, then mediation, then the courts of [venue] / good-faith "
                 "negotiation, then binding arbitration seated in [seat] under [rules]"), (T, ".")]),
    ("sub", [(B, "9.6 Whole understanding. "),
             (T, "This is the Partners’ entire understanding on its subject matter, replaces "
                 "prior discussions, and may be changed only in writing signed by every Partner.")]),

    ("p", [(B, "10. Signed.")]),
    ("sigtable", {
        "widths": [34, 44, 22],
        "head": ["Partner", "Signature", "Date"],
        "rows": [[[(F, "name")], [(T, "")], [(T, "")]],
                 [[(F, "name")], [(T, "")], [(T, "")]]],
    }),
]

FOOTER = ("Template — not legal advice. Review by qualified counsel before signature. "
          "Awake Venture Memorandum of Understanding v1.0 · CC BY 4.0 · cofounder.community/legal")


# --------------------------------------------------------------------------- Word

def build_docx(path):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.7)
    s.top_margin = s.bottom_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    def runs(par, parts, size=9.5):
        for kind, text in parts:
            r = par.add_run(text)
            r.font.size = Pt(size)
            if kind == B:
                r.bold = True
            elif kind == F:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.font.color.rgb = RGBColor(0x5A, 0x3A, 0x00)
        return par

    h = doc.add_paragraph()
    hr = h.add_run(TITLE.upper())
    hr.bold = True
    hr.font.size = Pt(15)
    h.paragraph_format.space_after = Pt(7)

    for line in PARTY:
        runs(doc.add_paragraph(), line)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_table(spec, sig=False):
        t = doc.add_table(rows=1, cols=len(spec["head"]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        usable = 8.5 - 1.4
        for i, w in enumerate(spec["widths"]):
            for cell in t.columns[i].cells:
                cell.width = Inches(usable * w / 100)
        for i, label in enumerate(spec["head"]):
            cell = t.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(label)
            r.bold = True
            r.font.size = Pt(8)
            r.font.name = "Calibri"
        for row in spec["rows"]:
            cells = t.add_row().cells
            for i, parts in enumerate(row):
                p = cells[i].paragraphs[0]
                for kind, text in parts:
                    r = p.add_run(text)
                    r.font.size = Pt(8)
                    r.font.name = "Calibri"
                    if kind == B:
                        r.bold = True
                    elif kind == F:
                        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                        r.font.color.rgb = RGBColor(0x5A, 0x3A, 0x00)
                if sig:
                    p.paragraph_format.space_before = Pt(7)
                    p.paragraph_format.space_after = Pt(7)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    for kind, payload in BODY:
        if kind == "p":
            runs(doc.add_paragraph(), payload)
        elif kind == "sub":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.space_after = Pt(3)
            runs(p, payload)
        elif kind == "table":
            add_table(payload)
        elif kind == "sigtable":
            add_table(payload, sig=True)

    foot = doc.sections[0].footer.paragraphs[0]
    fr = foot.add_run(FOOTER)
    fr.font.size = Pt(7)
    fr.font.name = "Calibri"
    fr.font.color.rgb = RGBColor(0x6D, 0x65, 0x5D)

    doc.save(path)
    return path


# --------------------------------------------------------------------------- HTML / PDF

def spans(parts):
    out = []
    for kind, text in parts:
        esc = html.escape(text)
        if kind == B:
            out.append(f"<strong>{esc}</strong>")
        elif kind == F:
            out.append(f'<span class="fill">{esc}</span>')
        else:
            out.append(esc)
    return "".join(out)


def build_html(path):
    css = """
@page { size: Letter; margin: 0.65in 0.7in 0.7in; }
* { box-sizing: border-box; }
body { margin: 0; color: #161411; font-family: Georgia, serif; font-size: 9.3pt; line-height: 1.32; }
main { max-width: 7.1in; margin: 0 auto; }
.doc-title { margin: 0 0 6pt; font-size: 15pt; font-weight: 700; letter-spacing: 0.02em; text-transform: uppercase; }
.party { margin: 0 0 8pt; padding-bottom: 7pt; border-bottom: 0.75pt solid #b83a2e; }
.party div { margin: 0 0 2pt; }
p { margin: 0 0 5pt; orphans: 3; widows: 3; }
p.sub { margin: 0 0 3pt; padding-left: 12pt; }
.fill { background: #fdf0b8; color: #5a3a00; padding: 0 2pt; border-radius: 1pt; }
table { width: 100%; margin: 4pt 0 7pt; border-collapse: collapse; table-layout: fixed; font-family: Calibri, Arial, sans-serif; font-size: 7.8pt; }
tr { break-inside: avoid; }
th, td { padding: 3pt 4pt; border: 0.5pt solid #cfc7bd; vertical-align: top; overflow-wrap: anywhere; }
th { background: #f2f0ec; text-align: left; font-weight: 700; }
.sigtable td { height: 24pt; }
footer { margin-top: 14pt; padding-top: 6pt; border-top: 0.5pt solid #cfc7bd; color: #6d655d;
         font-family: Calibri, Arial, sans-serif; font-size: 7pt; }
@media screen { body { background: #f2eee8; padding: 40px; } main { padding: 0.65in 0.7in; background: #fff; box-shadow: 0 8px 30px rgba(0,0,0,.12); } }
"""
    parts = [f'<!doctype html><html lang="en"><head><meta charset="utf-8">',
             f"<title>{TITLE} — fill-in copy</title><style>{css}</style></head><body><main>",
             f'<div class="doc-title">{html.escape(TITLE)}</div>', '<div class="party">']
    for line in PARTY:
        parts.append(f"<div>{spans(line)}</div>")
    parts.append("</div>")

    for kind, payload in BODY:
        if kind == "p":
            parts.append(f"<p>{spans(payload)}</p>")
        elif kind == "sub":
            parts.append(f'<p class="sub">{spans(payload)}</p>')
        elif kind in ("table", "sigtable"):
            cls = ' class="sigtable"' if kind == "sigtable" else ""
            parts.append(f"<table{cls}>")
            head = "".join(f'<th style="width:{w}%">{html.escape(h)}</th>'
                           for h, w in zip(payload["head"], payload["widths"]))
            parts.append(f"<tr>{head}</tr>")
            for row in payload["rows"]:
                parts.append("<tr>" + "".join(f"<td>{spans(c)}</td>" for c in row) + "</tr>")
            parts.append("</table>")

    parts.append(f"<footer>{html.escape(FOOTER)}</footer></main></body></html>")
    path.write_text("".join(parts), encoding="utf-8")
    return path


def main():
    docx_path = build_docx(HERE / f"{STEM}.docx")
    html_path = build_html(HERE / f"{STEM}.html")
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-sandbox",
                    "--no-pdf-header-footer", f"--print-to-pdf={HERE / f'{STEM}.pdf'}",
                    str(html_path)], capture_output=True)
    for p in (docx_path, html_path, HERE / f"{STEM}.pdf"):
        print(f"{p.name}: {p.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
